"""
ModernizeAI PowerPoint Presentation Generator
Creates an executive single-page widescreen (.pptx) presentation using python-pptx.
"""

import sys
import subprocess
from pathlib import Path

# Ensure python-pptx is available
try:
    import pptx
except ImportError:
    print("[INFO] python-pptx is not installed. Installing python-pptx...")
    try:
        subprocess.check_call([sys.executable, "-m", "pip", "install", "python-pptx"])
        import pptx
    except Exception as e:
        print(f"[ERROR] Could not install python-pptx automatically: {e}")
        print("Please run: pip install python-pptx")
        sys.exit(1)

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # Set 16:9 widescreen aspect ratio (13.333 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_slide_layout)

    # Color Palette (Dark Theme / Executive Cyberpunk)
    COLOR_BG = RGBColor(10, 15, 29)          # #0a0f1d
    COLOR_CARD = RGBColor(18, 26, 47)        # #121a2f
    COLOR_CARD_BORDER = RGBColor(56, 189, 248) # Cyan border
    COLOR_TITLE = RGBColor(56, 189, 248)     # #38bdf8
    COLOR_WHITE = RGBColor(248, 250, 252)    # #f8fafc
    COLOR_MUTED = RGBColor(148, 163, 184)    # #94a3b8
    COLOR_ACCENT_AMBER = RGBColor(245, 158, 11) # #f59e0b
    COLOR_ACCENT_PURPLE = RGBColor(192, 132, 252) # #c084fc
    COLOR_ACCENT_EMERALD = RGBColor(52, 211, 153) # #34d399
    COLOR_WARN_BG = RGBColor(40, 15, 25)     # Deep red background
    COLOR_WARN_BORDER = RGBColor(244, 63, 94) # Rose border
    COLOR_WARN_TEXT = RGBColor(254, 205, 211)

    # 1. Slide Background
    bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = COLOR_BG
    bg_shape.line.fill.background()

    # 2. Header Bar Background
    header_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.4), Inches(0.3), Inches(12.533), Inches(0.85))
    header_box.fill.solid()
    header_box.fill.fore_color.rgb = RGBColor(15, 23, 42)
    header_box.line.color.rgb = RGBColor(51, 65, 85)

    # Header Text
    tf_h = header_box.text_frame
    tf_h.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf_h.word_wrap = True
    p_h1 = tf_h.paragraphs[0]
    p_h1.text = "ModernizeAI: Enterprise Legacy Modernization Platform"
    p_h1.font.bold = True
    p_h1.font.size = Pt(17)
    p_h1.font.color.rgb = COLOR_TITLE

    p_h2 = tf_h.add_paragraph()
    p_h2.text = "Autonomous Hybrid Agentic Graph-RAG Swarm | Powered by Cognizant Neuro® AI Multi-Agent Accelerator (Neuro SAN Studio)"
    p_h2.font.size = Pt(10)
    p_h2.font.color.rgb = COLOR_MUTED

    # Helper function to add structured cards
    def add_card(x, y, w, h, title, items, title_color=COLOR_TITLE):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = COLOR_CARD_BORDER
        card.line.width = Pt(1)

        tf = card.text_frame
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.word_wrap = True
        tf.margin_left = Inches(0.12)
        tf.margin_right = Inches(0.12)
        tf.margin_top = Inches(0.10)
        tf.margin_bottom = Inches(0.10)

        # Title
        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.bold = True
        p_t.font.size = Pt(11)
        p_t.font.color.rgb = title_color
        p_t.space_after = Pt(4)

        # Content bullets
        for bold_prefix, text in items:
            p = tf.add_paragraph()
            p.space_after = Pt(3)
            if bold_prefix:
                run_bold = p.add_run()
                run_bold.text = bold_prefix + ": "
                run_bold.font.bold = True
                run_bold.font.size = Pt(9)
                run_bold.font.color.rgb = COLOR_WHITE

            run_text = p.add_run()
            run_text.text = text
            run_text.font.size = Pt(8.5)
            run_text.font.color.rgb = COLOR_MUTED

    # Row 1: 3 Columns
    # Card 1: Description & Purpose
    add_card(
        x=0.4, y=1.25, w=3.95, h=2.65,
        title="🎯 1. Description & Purpose",
        items=[
            ("Overview", "Enterprise-grade Agentic Knowledge Factory & Hybrid Graph-RAG platform solving legacy cloud migration risks."),
            ("The Problem", "Eliminates the 70%+ failure rate caused by unknown business rules, hidden database locks, and ripple effects."),
            ("Framework Origin", "Built atop Cognizant Neuro SAN Studio; coordinates an 11-agent mixture-of-experts swarm via HOCON & AAOSA protocol."),
            ("Key Scale", "11 Agents | 5 Coded Deterministic Tools | 5 Memory Tiers | 78/100 Readiness Score on ClaimCore v2.4 benchmark.")
        ]
    )

    # Card 2: Innovation & Core Ideas
    add_card(
        x=4.5, y=1.25, w=4.35, h=2.65,
        title="💡 2. Innovation & Core Ideas",
        items=[
            ("The 80/20 Rule", "80% Deterministic Extraction (AST code parsers, SQL DDL schema parsers, zero LLM hallucination) + 20% LLM Agent Reasoning."),
            ("5-Tier Memory Fabric", "Tier 1: Raw (SHA-256 + line citations) | Tier 2: Structural (AST symbols) | Tier 3: Semantic (Vector/BM25) | Tier 4: Procedural | Tier 5: Transformation."),
            ("Knowledge Graph", "Schema-enforced MultiDiGraph with Louvain community detection for microservice domain boundary isolation."),
            ("Line Provenance", "Every discovered entity is grounded directly in verifiable source code lines and database schemas.")
        ],
        title_color=COLOR_ACCENT_PURPLE
    )

    # Card 3: Key Features Delivered
    add_card(
        x=9.0, y=1.25, w=3.93, h=2.65,
        title="⚡ 3. Key Features Delivered",
        items=[
            ("Dual User Interfaces", "ModernizeAI Web App (PyVis 2D/3D interactive graph, blast radius explorer) & Neuro SAN Studio Client (nsflow topology)."),
            ("Business Rules Catalog", "Deterministically formalizes rules (BR-01 to BR-05) with exact line citations and implementing methods."),
            ("Blast-Radius Explorer", "Traces transitive upstream callers and downstream table impact for any proposed code/schema modification."),
            ("Sly-Data Privacy", "Safeguards proprietary code and schema credentials outside LLM context windows.")
        ],
        title_color=COLOR_ACCENT_EMERALD
    )

    # Row 2: Phase 2 Roadmap & Benchmark
    # Card 4: Future Upgrades (Phase 2 Swarm)
    add_card(
        x=0.4, y=4.0, w=8.45, h=2.0,
        title="🚀 4. Future Upgrades & Roadmap: Phase 2 Active Code Migration Swarm",
        items=[
            ("KG Reader Agent", "Ingests and slices candidate microservice subgraphs from the Knowledge Graph based on Louvain clustering."),
            ("Cloud Scaffolder Agent", "Scaffolds target cloud-native containerized microservices, Dockerfiles, and Helm/Terraform manifests."),
            ("Code Migrator Agent (.NET / Java)", "Automates transpilation: converts legacy on-prem .NET Framework (WCF/ADO.NET) to modern .NET 8/9 Minimal APIs / Cloud Run."),
            ("SQL to Cloud Decoupler & Test Synthesizer", "Decouples stored procedure row locks into Saga/Outbox patterns; synthesizes xUnit/JUnit tests directly from business rules for zero functional regression.")
        ],
        title_color=COLOR_ACCENT_AMBER
    )

    # Card 5: Benchmark Validation
    add_card(
        x=9.0, y=4.0, w=3.93, h=2.0,
        title="🏢 5. Legacy Benchmark Profile",
        items=[
            ("Application", "ClaimCore v2.4 (P&C Insurance Adjudication monolith)."),
            ("Tech Stack", "Java JDBC services, Oracle SQL schema, PL/SQL stored procedures (SP_PROCESS_CLAIM with pessimistic locks)."),
            ("Graph Topology", "21 verified nodes, 29 structural edges, 5 business rules."),
            ("Discrepancy Catch", "Discovered 15-day billing cutoff in SME tribal notes conflicting with 30-day documented spec rule.")
        ],
        title_color=COLOR_TITLE
    )

    # Row 3: Corporate Limitation Banner
    lim_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.4), Inches(6.1), Inches(12.533), Inches(0.95))
    lim_box.fill.solid()
    lim_box.fill.fore_color.rgb = COLOR_WARN_BG
    lim_box.line.color.rgb = COLOR_WARN_BORDER
    lim_box.line.width = Pt(1.5)

    tf_l = lim_box.text_frame
    tf_l.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf_l.word_wrap = True
    tf_l.margin_left = Inches(0.15)
    tf_l.margin_right = Inches(0.15)

    p_lw = tf_l.paragraphs[0]
    r_badge = p_lw.add_run()
    r_badge.text = "⚠️ PRODUCTION & CORPORATE ENVIRONMENT LIMITATION (COGNIZANT LAPTOP RESTRICTION): "
    r_badge.font.bold = True
    r_badge.font.size = Pt(9.5)
    r_badge.font.color.rgb = COLOR_WARN_BORDER

    p_lt = tf_l.add_paragraph()
    r_body = p_lt.add_run()
    r_body.text = (
        "ModernizeAI relies on active external LLM API keys (OpenAI, Anthropic, or Google Gemini) for its 20% LLM reasoning swarm. "
        "In the current corporate environment on Cognizant laptops, corporate network security (Zscaler / Scalar) blocks all outbound Generative AI API calls and LLM endpoints. "
        "Consequently, live end-to-end multi-agent chat and LLM generation cannot be executed or demoed directly on corporate PCs without an approved cloud sandbox or network proxy exemption."
    )
    r_body.font.size = Pt(8.5)
    r_body.font.color.rgb = COLOR_WARN_TEXT

    # 4. Slide Footer
    footer_box = slide.shapes.add_textbox(Inches(0.4), Inches(7.1), Inches(12.533), Inches(0.3))
    tf_f = footer_box.text_frame
    p_f = tf_f.paragraphs[0]
    p_f.text = "ModernizeAI | Cognizant AI Lab • Neuro SAN Studio Accelerator | Single-Page Executive Presentation (.pptx)"
    p_f.font.size = Pt(8)
    p_f.font.color.rgb = RGBColor(100, 116, 139)

    output_path = Path(__file__).resolve().parent.parent / "artifacts" / "ModernizeAI_Executive_Presentation.pptx"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
    print(f"[SUCCESS] PowerPoint presentation saved to: {output_path}")

if __name__ == "__main__":
    create_presentation()
